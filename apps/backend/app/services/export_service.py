import io
from datetime import datetime, timezone

from app.models.entities import Workspace


def export_workspace_markdown(ws: Workspace) -> str:
    lines = []
    lines.append(f"# Workspace: {ws.name}")
    if ws.description:
        lines.append(f"**Description:** {ws.description}\n")
    lines.append(f"**Exported:** {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}")
    lines.append(f"**Created:** {ws.created_at.strftime('%Y-%m-%d %H:%M:%S UTC') if ws.created_at else 'N/A'}\n")
    
    # Summary
    lines.append("## Overview")
    lines.append(f"- **Notebooks / Canvases:** {len(ws.canvas_documents)}")
    lines.append(f"- **Notes:** {len(ws.notes)}")
    lines.append(f"- **Conversations:** {len(ws.conversations)}")
    lines.append(f"- **Files Attached:** {len(ws.files)}\n")
    lines.append("---\n")
    
    # Notebooks & Canvas
    lines.append("## Notebooks & Canvases\n")
    if not ws.canvas_documents:
        lines.append("*No notebooks found in this workspace.*\n")
    else:
        for idx, doc in enumerate(ws.canvas_documents, 1):
            lines.append(f"### {idx}. {doc.name}")
            lines.append(f"- **Version:** {doc.document_version}")
            lines.append(f"- **Last Updated:** {doc.updated_at.strftime('%Y-%m-%d %H:%M') if doc.updated_at else 'N/A'}")
            
            # Extract any text objects from canvas
            text_objects = [obj for obj in doc.objects if obj.object_type == "text"]
            if text_objects:
                lines.append("\n**Text Content:**")
                for tobj in text_objects:
                    text_val = (tobj.properties or {}).get("text") or (tobj.bounds or {}).get("text") or ""
                    if text_val:
                        lines.append(f"> {text_val}")
            lines.append("")
            
    # Notes
    lines.append("---\n## Notes & Documents\n")
    if not ws.notes:
        lines.append("*No standalone notes found.*\n")
    else:
        for idx, note in enumerate(ws.notes, 1):
            lines.append(f"### {idx}. {note.title}")
            lines.append(f"- **Date:** {note.updated_at.strftime('%Y-%m-%d %H:%M') if note.updated_at else 'N/A'}\n")
            lines.append(note.content or "*Empty note content*")
            lines.append("")

    # Conversations
    lines.append("---\n## Conversations & AI Chats\n")
    if not ws.conversations:
        lines.append("*No conversation threads found.*\n")
    else:
        for idx, conv in enumerate(ws.conversations, 1):
            lines.append(f"### {idx}. {conv.title}")
            lines.append(f"- **Status:** {conv.status}")
            lines.append(f"- **Date:** {conv.updated_at.strftime('%Y-%m-%d %H:%M') if conv.updated_at else 'N/A'}\n")
            if not conv.messages:
                lines.append("*No messages in this chat.*\n")
            else:
                for msg in conv.messages:
                    speaker = "**User**" if msg.role == "user" else "**NoteFlow AI**"
                    lines.append(f"#### {speaker}:")
                    lines.append(f"{msg.content}\n")
                    
    # Files
    if ws.files:
        lines.append("---\n## Attached Files\n")
        for f in ws.files:
            lines.append(f"- **{f.filename}** ({f.file_type or 'unknown'}, {round((f.file_size or 0) / 1024, 1)} KB)")
        lines.append("")

    return "\n".join(lines)


def export_workspace_docx(ws: Workspace) -> bytes:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    from docx.shared import Pt, RGBColor

    doc = Document()

    # Document Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(ws.name)
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(16, 185, 129) # Emerald Brand Color

    if ws.description:
        desc_p = doc.add_paragraph()
        desc_run = desc_p.add_run(ws.description)
        desc_run.font.size = Pt(12)
        desc_run.font.italic = True
        desc_run.font.color.rgb = RGBColor(100, 116, 139)

    meta_p = doc.add_paragraph()
    meta_run = meta_p.add_run(
        f"Exported on {datetime.now(timezone.utc).strftime('%B %d, %Y at %H:%M UTC')} | NoteFlow Workspace"
    )
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = RGBColor(148, 163, 184)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Summary Table
    table = doc.add_table(rows=2, cols=4)
    table.style = 'Table Grid'
    headers = ['Notebooks', 'Notes', 'Chats', 'Files']
    counts = [str(len(ws.canvas_documents)), str(len(ws.notes)), str(len(ws.conversations)), str(len(ws.files))]
    
    for i, h in enumerate(headers):
        hdr_cell = table.cell(0, i)
        hdr_cell.text = h
        hdr_cell.paragraphs[0].runs[0].font.bold = True
        hdr_cell.paragraphs[0].runs[0].font.size = Pt(10)
        # Light grey background
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
        hdr_cell._tc.get_or_add_tcPr().append(shd)

        val_cell = table.cell(1, i)
        val_cell.text = counts[i]
        val_cell.paragraphs[0].runs[0].font.bold = True
        val_cell.paragraphs[0].runs[0].font.size = Pt(13)
        val_cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(37, 99, 235)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # 1. Notebooks & Canvases
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run('1. Notebooks & Canvases')
    h1_run.font.color.rgb = RGBColor(15, 23, 42)

    if not ws.canvas_documents:
        p = doc.add_paragraph('No notebooks currently in this workspace.')
        p.runs[0].font.italic = True
    else:
        for idx, cdoc in enumerate(ws.canvas_documents, 1):
            h2 = doc.add_heading(level=2)
            h2.add_run(f'{idx}. {cdoc.name}')
            
            p_info = doc.add_paragraph()
            p_info.add_run(f"Objects Count: {len(cdoc.objects)} | Version: {cdoc.document_version}").font.size = Pt(9.5)
            
            text_objects = [obj for obj in cdoc.objects if obj.object_type == "text"]
            if text_objects:
                p_sub = doc.add_paragraph()
                p_sub.add_run("Text & Notes:").font.bold = True
                for tobj in text_objects:
                    text_val = (tobj.properties or {}).get("text") or (tobj.bounds or {}).get("text") or ""
                    if text_val:
                        p_t = doc.add_paragraph(style='Quote')
                        p_t.add_run(text_val)

    # 2. Standalone Notes
    h1_notes = doc.add_heading(level=1)
    h1_notes.add_run('2. Notes & Documentation')
    h1_notes.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    if not ws.notes:
        p = doc.add_paragraph('No standalone notes in this workspace.')
        p.runs[0].font.italic = True
    else:
        for idx, note in enumerate(ws.notes, 1):
            h2 = doc.add_heading(level=2)
            h2.add_run(f'{idx}. {note.title}')
            p = doc.add_paragraph(note.content or '')
            p.paragraph_format.space_after = Pt(8)

    # 3. Conversations & AI Chats
    h1_chats = doc.add_heading(level=1)
    h1_chats.add_run('3. Conversations & AI Chats')
    h1_chats.runs[0].font.color.rgb = RGBColor(15, 23, 42)

    if not ws.conversations:
        p = doc.add_paragraph('No conversation records.')
        p.runs[0].font.italic = True
    else:
        for idx, conv in enumerate(ws.conversations, 1):
            h2 = doc.add_heading(level=2)
            h2.add_run(f'{idx}. {conv.title}')
            for msg in conv.messages:
                p = doc.add_paragraph()
                is_user = msg.role == 'user'
                speaker_run = p.add_run("User: " if is_user else "NoteFlow AI: ")
                speaker_run.font.bold = True
                speaker_run.font.color.rgb = RGBColor(30, 41, 59) if is_user else RGBColor(16, 185, 129)
                p.add_run(msg.content or '')
                p.paragraph_format.space_after = Pt(6)

    # Save to buffer
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_workspace_pdf(ws: Workspace) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'WorkspaceTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#0F172A'),
        alignment=0, # Left-aligned
        spaceAfter=6,
    )
    
    subtitle_style = ParagraphStyle(
        'WorkspaceSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=10,
    )
    
    meta_style = ParagraphStyle(
        'WorkspaceMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#94A3B8'),
        spaceAfter=14,
    )
    
    h1_style = ParagraphStyle(
        'Heading1Style',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=14,
        spaceAfter=8,
    )
    
    h2_style = ParagraphStyle(
        'Heading2Style',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#2563EB'),
        spaceBefore=10,
        spaceAfter=4,
    )
    
    body_style = ParagraphStyle(
        'BodyStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6,
    )
    
    quote_style = ParagraphStyle(
        'QuoteStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#475569'),
        leftIndent=15,
        spaceAfter=6,
    )

    user_msg_style = ParagraphStyle(
        'UserMsgStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=4,
    )

    ai_msg_style = ParagraphStyle(
        'AiMsgStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#0F766E'),
        leftIndent=10,
        spaceAfter=8,
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph(ws.name, title_style))
    if ws.description:
        story.append(Paragraph(ws.description, subtitle_style))
    
    date_str = datetime.now(timezone.utc).strftime('%B %d, %Y - %H:%M UTC')
    story.append(Paragraph(f"Workspace Export Document &bull; Generated on {date_str}", meta_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#10B981'), spaceAfter=14))

    # Metric Table
    metric_data = [
        ['Notebooks', 'Notes', 'Conversations', 'Attached Files'],
        [str(len(ws.canvas_documents)), str(len(ws.notes)), str(len(ws.conversations)), str(len(ws.files))]
    ]
    t = Table(metric_data, colWidths=[130, 130, 130, 130])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F8FAFC')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#475569')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, 1), 14),
        ('TEXTCOLOR', (0, 1), (-1, 1), colors.HexColor('#2563EB')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 1), (-1, 1), 8),
        ('BOTTOMPADDING', (0, 1), (-1, 1), 8),
    ]))
    story.append(t)
    story.append(Spacer(1, 16))

    # Section 1: Notebooks & Canvases
    story.append(Paragraph("1. Notebooks & Canvases", h1_style))
    if not ws.canvas_documents:
        story.append(Paragraph("<i>No notebooks in this workspace.</i>", body_style))
    else:
        for idx, cdoc in enumerate(ws.canvas_documents, 1):
            story.append(Paragraph(f"{idx}. {cdoc.name}", h2_style))
            story.append(Paragraph(f"Objects Count: {len(cdoc.objects)} | Version: {cdoc.document_version}", meta_style))
            
            text_objects = [obj for obj in cdoc.objects if obj.object_type == "text"]
            for tobj in text_objects:
                text_val = (tobj.properties or {}).get("text") or (tobj.bounds or {}).get("text") or ""
                if text_val:
                    # Sanitize for reportlab XML
                    clean_text = text_val.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    story.append(Paragraph(f"&bull; {clean_text}", quote_style))
    story.append(Spacer(1, 14))

    # Section 2: Notes
    story.append(Paragraph("2. Standalone Notes", h1_style))
    if not ws.notes:
        story.append(Paragraph("<i>No standalone notes in this workspace.</i>", body_style))
    else:
        for idx, note in enumerate(ws.notes, 1):
            story.append(Paragraph(f"{idx}. {note.title}", h2_style))
            content_safe = (note.content or '').replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            story.append(Paragraph(content_safe, body_style))
    story.append(Spacer(1, 14))

    # Section 3: Conversations
    story.append(Paragraph("3. Conversations & AI Chats", h1_style))
    if not ws.conversations:
        story.append(Paragraph("<i>No conversations in this workspace.</i>", body_style))
    else:
        for idx, conv in enumerate(ws.conversations, 1):
            story.append(Paragraph(f"{idx}. {conv.title}", h2_style))
            for msg in conv.messages:
                content_safe = (msg.content or '').replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                if msg.role == 'user':
                    story.append(Paragraph(f"<b>User:</b> {content_safe}", user_msg_style))
                else:
                    story.append(Paragraph(f"<b>NoteFlow AI:</b> {content_safe}", ai_msg_style))

    pdf.build(story)
    return buf.getvalue()
